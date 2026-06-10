"""
命令列匯入靈動商品總表到 Firestore（支援多分頁），與網頁 import.js 邏輯等價。

網頁的「匯入產品總表」按鈕一次只能讀第一個分頁；本工具會掃描所有分頁、
合併去重（以分流碼為唯一鍵），並以 merge:true upsert 到 Products。
merge 模式會保留既有的 imageUrl / imageSource 等雲端欄位，只覆蓋總表提供的欄位。

用法：
  python tools/sync_quote_to_firestore.py --input "<xlsx>"            # dry-run 預覽差異
  python tools/sync_quote_to_firestore.py --input "<xlsx>" --commit   # 實際寫入
  python tools/sync_quote_to_firestore.py --input "<xlsx>" --prefix 515  # 只處理某前綴
"""
import argparse
import datetime
import json
import math
import os
import re
import sys
from collections import Counter

from openpyxl import load_workbook
import firebase_admin
from firebase_admin import credentials, firestore


def nkey(v):
    return re.sub(r"\s+", "", str(v or "")).replace("\r", "").replace("\n", "").strip()


def to_text(v):
    if v is None:
        return ""
    t = str(v).strip()
    return "" if t.lower() == "nan" else t


def to_number(v):
    t = to_text(v).replace(",", "")
    if not t:
        return 0
    try:
        return round(float(t), 2)
    except ValueError:
        return 0


def to_barcode(v):
    t = to_text(v).replace(",", "")
    if not t:
        return ""
    return t.split(".")[0] if re.match(r"^\d+\.0+$", t) else t


def to_status(s):
    s = to_text(s)
    return "inactive" if s in ("下架", "停產") else "active"


def to_inventory(s):
    return 0 if to_text(s) == "缺貨中" else 200


def norm_code(v):
    if v is None:
        return ""
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    t = str(v).strip()
    return t.split(".")[0] if re.match(r"^\d+\.0+$", t) else t


DIVISOR_MAP = {
    4: {50: 0.75, 100: 0.78, 300: 0.81, 500: 0.835, 1000: 0.858, 3000: 0.89},
    3: {50: 0.75, 100: 0.78, 300: 0.81, 500: 0.835, 1000: 0.858},
    2: {50: 0.74, 100: 0.77, 300: 0.80},
    1: {50: 0.73, 100: 0.76},
}


def quotes_for(cost):
    """各等級各級距的預存報價 {level: {tier: price}}（方案 B：前端讀此值、不需 cost）。"""
    try:
        c = float(cost)
    except (TypeError, ValueError):
        return {}
    if c <= 0:
        return {}
    return {
        str(lv): {str(t): int(math.ceil((c / d) * 1.05)) for t, d in tiers.items()}
        for lv, tiers in DIVISOR_MAP.items()
    }


def find_header(rows):
    for r in range(min(len(rows), 20)):
        m = {}
        for idx, cell in enumerate(rows[r] or []):
            k = nkey(cell)
            if k:
                m[k] = idx
        if "分流" in m:
            return r, m
    return -1, {}


def parse_workbook(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    items = {}
    sheet_stats = {}
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        hidx, hmap = find_header(rows)
        if hidx == -1:
            sheet_stats[sn] = 0
            continue

        def gv(row, name, hmap=hmap):
            if name not in hmap:
                return None
            i = hmap[name]
            return row[i] if i < len(row) else None

        cnt = 0
        for r in range(hidx + 1, len(rows)):
            row = rows[r] or []
            sc = norm_code(gv(row, "分流"))
            if not sc:
                continue
            model = to_text(gv(row, "型號")).upper() or sc
            st = to_text(gv(row, "狀態"))
            items[sc] = dict(
                splitCode=sc,
                brand=to_text(gv(row, "品牌")),
                category=to_text(gv(row, "分類")),
                model=model,
                mainModel=model.split("-")[0],
                name=to_text(gv(row, "商品名稱")) or model,
                cost=to_number(gv(row, "詢價含")),
                marketPrice=to_number(gv(row, "市價含")),
                minPrice=to_number(gv(row, "售價含")),
                inventory=to_inventory(st),
                status=to_status(st),
                statusText=st,
                isControlled=False,
                internationalBarcode=to_barcode(gv(row, "國際條碼")),
                barcode=to_barcode(gv(row, "國際條碼")),
                cartonQty=int(to_number(gv(row, "箱入數")) or 0),
                bsmi=to_text(gv(row, "BSMI")),
                ncc=to_text(gv(row, "NCC")),
                productUrl=to_text(gv(row, "商品對應網站")),
                netSalesPermission="",
                sourceSheet=sn,
            )
            # 總表慣例：價格 0/空 = 未提供（如「最低特價含」整欄 0）→ 不寫入，merge 保留線上既有值
            for pf in ("cost", "marketPrice", "minPrice"):
                if not items[sc].get(pf):
                    items[sc].pop(pf, None)
            q = quotes_for(items[sc].get("cost"))
            if q:  # cost 未提供時不寫 quotes，避免空 {} 蓋掉線上預存報價
                items[sc]["quotes"] = q
            cnt += 1
        sheet_stats[sn] = cnt
    return items, sheet_stats


def get_db():
    sa = os.path.join(os.path.dirname(__file__), "..", "backend", "serviceAccountKey.json")
    cred = credentials.Certificate(sa)
    if not firebase_admin._apps:
        firebase_admin.initialize_app(cred)
    return firestore.client()


CMP_FIELDS = ["cost", "marketPrice", "minPrice", "status", "statusText", "name"]
PRICE_FIELDS = {"cost", "marketPrice", "minPrice"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--commit", action="store_true")
    ap.add_argument("--prefix", default="", help="只處理分流碼以此前綴開頭的商品")
    args = ap.parse_args()

    items, sheet_stats = parse_workbook(args.input)
    if args.prefix:
        items = {k: v for k, v in items.items() if str(k).startswith(args.prefix)}

    db = get_db()
    online = {d.id: d.to_dict() for d in db.collection("Products").stream()}

    new, changed, same = [], [], []
    for sc, obj in items.items():
        if sc not in online:
            new.append(obj)
            continue
        cur = online[sc]
        diffs = []
        for f in CMP_FIELDS:
            if f in PRICE_FIELDS and f not in obj:
                continue  # 價格未提供（0/空）→ 不寫入也不列入差異
            ov, nv = cur.get(f), obj.get(f)
            if f in PRICE_FIELDS:
                if round(float(ov or 0), 2) != round(float(nv or 0), 2):
                    diffs.append((f, ov, nv))
            elif str(ov or "") != str(nv or ""):
                diffs.append((f, ov, nv))
        (changed if diffs else same).append((obj, diffs) if diffs else obj)

    print(f"解析: {sum(sheet_stats.values())} 列 -> 去重 {len(items)} 筆")
    print("各分頁:", dict(sheet_stats))
    print(f"線上現有: {len(online)} 筆")
    print(f"\n新增(新品): {len(new)} 筆")
    print(f"更新(價格/狀態/名稱有變): {len(changed)} 筆")
    print(f"不變(完全相同): {len(same)} 筆")

    print(f"\n=== 將更新的重複項 {len(changed)} 筆（線上 -> 新總表）===")
    for obj, diffs in sorted(changed, key=lambda x: x[0]["splitCode"]):
        ds = ", ".join(f"{f}:{ov}->{nv}" for f, ov, nv in diffs)
        print(f'  {obj["splitCode"]:8} {obj["model"]:14} | {ds}')

    if not args.commit:
        print("\n[DRY-RUN] 未寫入任何資料。確認無誤後加 --commit 實際寫入。")
        return

    # 寫入前自動備份當前線上 Products（安全網，可還原）
    bdir = os.path.join(os.path.dirname(__file__), "..", "backups")
    os.makedirs(bdir, exist_ok=True)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    bpath = os.path.join(bdir, f"Products_{ts}.json")
    with open(bpath, "w", encoding="utf-8") as f:
        json.dump([{"_id": k, **v} for k, v in online.items()], f,
                  ensure_ascii=False, indent=2, default=str)
    print(f"\n[BACKUP] 已備份線上 {len(online)} 筆 Products -> {bpath}")

    batch = db.batch()
    n = total = 0
    for sc, obj in items.items():
        batch.set(db.collection("Products").document(sc), obj, merge=True)
        n += 1
        total += 1
        if n >= 400:
            batch.commit()
            batch = db.batch()
            n = 0
    if n > 0:
        batch.commit()
    print(f"\n[COMMIT] 已 merge upsert {total} 筆到 Products。")


if __name__ == "__main__":
    main()
